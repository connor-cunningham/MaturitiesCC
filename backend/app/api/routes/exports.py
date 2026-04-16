from datetime import date
from typing import Optional
from io import BytesIO
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.core.database import get_db
from app.models.canonical import CanonicalLoan, CanonicalProperty, CanonicalOwner
from app.models.crm import OutreachTarget

router = APIRouter(prefix="/exports", tags=["exports"])


def _excel_response(wb: Workbook, filename: str) -> StreamingResponse:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _style_header(ws):
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


@router.get("/loans")
async def export_loans(
    maturity_start: Optional[date] = None,
    maturity_end: Optional[date] = None,
    state: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    query = select(CanonicalLoan).options(
        selectinload(CanonicalLoan.property),
        selectinload(CanonicalLoan.owner),
    )
    filters = []
    if maturity_start:
        filters.append(CanonicalLoan.maturity_date >= maturity_start)
    if maturity_end:
        filters.append(CanonicalLoan.maturity_date <= maturity_end)
    if state:
        query = query.join(CanonicalProperty, CanonicalLoan.property_id == CanonicalProperty.id)
        filters.append(CanonicalProperty.state == state.upper())
    if filters:
        query = query.where(and_(*filters))
    query = query.order_by(CanonicalLoan.maturity_date.asc())
    result = await db.execute(query)
    loans = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Loans"
    headers = ["Property", "City", "State", "Owner", "Lender", "Originator", "Servicer",
               "Origination Date", "Maturity Date", "Months to Maturity",
               "Original Amount", "Current Balance", "Rate Type", "Coupon",
               "IO Flag", "Loan Type", "Recourse", "Priority Score"]
    ws.append(headers)

    today = date.today()
    for l in loans:
        months = round((l.maturity_date - today).days / 30.44, 1) if l.maturity_date else None
        ws.append([
            l.property.display_name if l.property else "",
            l.property.city if l.property else "",
            l.property.state if l.property else "",
            l.owner.display_name if l.owner else "",
            l.lender or "", l.originator or "", l.servicer or "",
            str(l.origination_date) if l.origination_date else "",
            str(l.maturity_date) if l.maturity_date else "",
            months,
            l.original_amount, l.current_balance,
            l.rate_type or "", l.coupon,
            "Yes" if l.io_flag else "No" if l.io_flag is not None else "",
            l.loan_type or "",
            "Yes" if l.recourse else "No" if l.recourse is not None else "",
            l.priority_score,
        ])

    _style_header(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    return _excel_response(wb, "loans_export.xlsx")


@router.get("/properties")
async def export_properties(
    state: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    query = select(CanonicalProperty).options(selectinload(CanonicalProperty.owner))
    if state:
        query = query.where(CanonicalProperty.state == state.upper())
    result = await db.execute(query)
    props = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Properties"
    ws.append(["Name", "Street", "City", "State", "Zip", "Submarket",
               "Owner", "Units", "Year Built", "Class", "Type",
               "Occupancy", "Vacancy", "Priority Score"])
    for p in props:
        ws.append([
            p.display_name, p.street or "", p.city or "", p.state or "", p.zip or "",
            p.submarket or "", p.owner.display_name if p.owner else "",
            p.units, p.year_built, p.building_class or "", p.property_type or "",
            p.occupancy, p.vacancy, p.priority_score,
        ])

    _style_header(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    return _excel_response(wb, "properties_export.xlsx")


@router.get("/owners")
async def export_owners(db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(CanonicalOwner).options(
            selectinload(CanonicalOwner.properties),
            selectinload(CanonicalOwner.loans),
        )
    )
    owners = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Owners"
    today = date.today()
    ws.append(["Owner", "HQ City", "HQ State", "Outreach Stage", "Target Tier",
               "Properties", "Loans", "Total Units", "Upcoming Maturities (12mo)",
               "Total Upcoming Volume", "Last Contact", "Next Follow-up", "Priority Score"])
    for o in owners:
        loans = o.loans or []
        upcoming_12 = [l for l in loans if l.maturity_date and (l.maturity_date - today).days <= 365]
        ws.append([
            o.display_name, o.hq_city or "", o.hq_state or "",
            o.outreach_stage or "", o.target_tier or "",
            len(o.properties or []), len(loans),
            sum(p.units or 0 for p in (o.properties or [])),
            len(upcoming_12),
            sum(l.original_amount or 0 for l in upcoming_12),
            str(o.last_contact_date) if o.last_contact_date else "",
            str(o.next_followup_date) if o.next_followup_date else "",
            o.priority_score,
        ])

    _style_header(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    return _excel_response(wb, "owners_export.xlsx")


@router.get("/targets")
async def export_targets(db: AsyncSession = Depends(get_db)):
    result = await db.execute(
        select(OutreachTarget).order_by(OutreachTarget.updated_at.desc())
    )
    targets = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"
    ws.append(["Owner ID", "Property ID", "Stage", "Tier",
               "Last Contact", "Next Follow-up", "Relationship", "Tags", "Notes"])
    for t in targets:
        ws.append([
            str(t.owner_id) if t.owner_id else "",
            str(t.property_id) if t.property_id else "",
            t.stage, t.target_tier or "",
            str(t.last_contact_date) if t.last_contact_date else "",
            str(t.next_followup_date) if t.next_followup_date else "",
            t.relationship_strength or "",
            ", ".join(t.tags or []),
            t.internal_notes or "",
        ])

    _style_header(ws)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    return _excel_response(wb, "targets_export.xlsx")
